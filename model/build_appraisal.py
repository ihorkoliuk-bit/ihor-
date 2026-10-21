import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

OUT = sys.argv[1]
wb = Workbook()
F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="404040")
YEL = PatternFill("solid", fgColor="FFFF00")
SEC = PatternFill("solid", fgColor="D9D9D9")
GBP = '£#,##0;(£#,##0);-'
PCT = '0.0%;(0.0%);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'
thin = Side(style="thin", color="999999")

def hdr(ws, row, labels):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l)
        c.font = HDR; c.fill = HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")

def section(ws, row, text, ncol=4):
    for col in range(1, ncol + 1):
        ws.cell(row=row, column=col).fill = SEC
    ws.cell(row=row, column=1, value=text).font = BOLD

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Припущення"
P = "'Припущення'!"
inp["A1"] = "ZON Wivenhoe (Croquet Gardens) — припущення моделі"; inp["A1"].font = TITLE
inp["A2"] = ("Легенда: синій текст = вхідні дані (можна змінювати); чорний = формули; "
             "зелений = посилання на інший аркуш; жовта заливка = ключове припущення, "
             "яке треба замінити реальними даними перед показом інвесторам.")
inp["A2"].font = Font(name=F, italic=True); inp.merge_cells("A2:K2")
inp["A2"].alignment = Alignment(wrap_text=True); inp.row_dimensions[2].height = 30

hdr(inp, 4, ["Тип юніта", "GIA, м²", "Спальні", "Мешканців", "Поверхів",
             "NDSS мін., м²", "NDSS ок?", "Ціна продажу, £/м²", "Ціна юніта, £",
             "Ділянка (сад+будинок), м²", "Джерело / примітка"])
types = [
    ("Lake House Compact 61", 61.5, 2, 3, 1, 61, "=340000/B{r}", 340,
     "GIA з брифу. NDSS 2b3p 1 пов. = 61 м². Ціна £340k — робоча ціна команди ZON, не підтверджена агентом."),
    ("Meadow 71", 71, 2, 4, 1, 70, "=390000/B{r}", 370,
     "GIA з брифу. NDSS 2b4p 1 пов. = 70 м². Ціна £390k — робоча ціна команди ZON, не підтверджена агентом."),
    ("Country 140", 140, 3, 6, 1, 95, "=550000/B{r}", 560,
     "GIA з брифу. NDSS 3b6p 1 пов. = 95 м². Ціна £550k — робоча ціна команди ZON, не підтверджена агентом."),
]
TYPE_ROWS = []
for i, t in enumerate(types):
    r = 5 + i; TYPE_ROWS.append(r)
    name, gia, beds, ppl, st, ndss, price, plot, note = t
    inp.cell(r, 1, name).font = BLACK
    for col, v in zip([2, 3, 4, 5, 6, 8, 10], [gia, beds, ppl, st, ndss, price, plot]):
        if isinstance(v, str): v = v.format(r=r)
        c = inp.cell(r, col, v); c.font = BLUE
    inp.cell(r, 7, f'=IF(B{r}>=F{r},"Так","НІ")').font = BLACK
    inp.cell(r, 9, f"=B{r}*H{r}").font = BLACK
    inp.cell(r, 11, note).font = BLACK
    inp.cell(r, 2).number_format = NUM1
    inp.cell(r, 8).number_format = GBP
    inp.cell(r, 9).number_format = GBP
    inp.cell(r, 10).number_format = NUM
    inp.cell(r, 8).fill = YEL
inp["A9"] = ("Ціни — робочі ціни команди ZON (£340k / £390k / £550k). "
             "Треба підтвердити звітом агента (Savills / Fenn Wright / Beresfords).")
inp["A9"].font = Font(name=F, italic=True); inp.merge_cells("A9:K9")

A = {}  # name -> absolute ref
def put(row, key, label, value, unit, note, fmt, yellow=False):
    inp.cell(row, 1, label).font = BLACK
    c = inp.cell(row, 2, value); c.font = BLUE; c.number_format = fmt
    if yellow: c.fill = YEL
    inp.cell(row, 3, unit).font = BLACK
    inp.cell(row, 4, note).font = BLACK
    A[key] = f"{P}$B${row}"

r = 11
section(inp, r, "Ділянка та планування"); r += 1
put(r, "site", "Площа ділянки", 14245, "м²", "Від користувача: 3,52 акра", NUM); r += 1
put(r, "roads", "Дороги, тротуари, паркування", 0.18, "% ділянки", "Припущення: петля Manual for Streets + гостьові місця", PCT); r += 1
put(r, "suds", "SuDS (свейли, атенюація) та відкритий простір", 0.07, "% ділянки", "Без струмка. Припущення; уточнює інженер з дренажу після тестів інфільтрації", PCT); r += 1
put(r, "bng", "BNG / ландшафтні буфери на ділянці", 0.06, "% ділянки", "Припущення; уточнює еколог (метрика DEFRA)", PCT); r += 1
put(r, "hubland", "Hub: забудова + тераса + паркування", 900, "м²", "Припущення: 175 м² пляма + тераса + ~10 місць", NUM); r += 1
r += 1
section(inp, r, "Доступне житло (S106)"); r += 1
put(r, "aff", "Частка доступного житла", 0.30, "% юнітів", "Припущення за політикою Colchester; перевірити в угоді S106", PCT, True); r += 1
put(r, "affval", "Ціна продажу RSL", 0.60, "% ринкової", "Припущення: мікс affordable rent + shared ownership", PCT, True); r += 1
put(r, "afftype", "Який тип іде як доступне житло", 1, "1/2/3", "1 = Compact 61, 2 = Meadow 71, 3 = Country 140", NUM); r += 1
r += 1
section(inp, r, "Вартість будівництва"); r += 1
put(r, "bc", "Будівництво бунгало (під ключ)", "='Собівартість будинку'!$E$27", "£/м² GIA", "Розрахунок з оболонки ZON £1 700/м² — аркуш «Собівартість будинку»", GBP, True); inp.cell(r, 2).font = GREEN; r += 1
put(r, "bchub", "Будівництво ZON Hub", 2800, "£/м² GIA", "Припущення: змішане використання, комерційне обладнання", GBP); r += 1
put(r, "ext", "Зовнішні роботи та мережі", 0.15, "% будівництва", "Припущення: дороги, SuDS, містки, мережі, підстанція", PCT); r += 1
put(r, "cont", "Резерв на непередбачене", 0.05, "% буд.+зовн.", "Припущення", PCT); r += 1
put(r, "fees", "Проєктування та консультанти", 0.10, "% буд.+зовн.", "Припущення", PCT); r += 1
put(r, "s106", "Внески S106 / інфраструктура", 15000, "£/юніт", "Припущення; уточнити з угоди S106", GBP); r += 1
put(r, "rams", "Essex Coast RAMS", 160, "£/юніт", "Приблизний тариф; перевірити поточну ставку", GBP); r += 1
put(r, "bngoff", "BNG за межами ділянки (кредити)", 50000, "£ разом", "Припущення", GBP); r += 1
r += 1
section(inp, r, "Продаж і фінансування"); r += 1
put(r, "sales", "Маркетинг і агенти", 0.025, "% GDV", "Припущення", PCT); r += 1
put(r, "legal", "Юридичні витрати на продаж", 1500, "£/юніт", "Припущення", GBP); r += 1
put(r, "rate", "Ставка девелоперського кредиту", 0.09, "% річних", "Припущення", PCT); r += 1
put(r, "draw", "Середня вибірка кредиту", 0.50, "% витрат", "Стандартне спрощення S-кривої", PCT); r += 1
put(r, "months", "Програма: будівництво + продаж", 24, "місяців", "Припущення", NUM); r += 1
r += 1
section(inp, r, "Земля і ціль"); r += 1
put(r, "land", "Ціна землі", 1500000, "£", "ЗАГЛУШКА — замінити фактичною ціною / ціною опціону", GBP, True); r += 1
put(r, "landcost", "SDLT + юридичні на землю", 0.06, "% ціни землі", "Припущення (ставки для нежитлової нерухомості)", PCT); r += 1
put(r, "target", "Цільова маржа девелопера", 0.20, "% GDV", "Типовий орієнтир кредиторів 17–20%", PCT); r += 1
r += 1
section(inp, r, "ZON Hub (лише сценарій B)"); r += 1
put(r, "hubgia", "Hub GIA разом", 350, "м²", "Бриф: 175 + 175 м²", NUM); r += 1
put(r, "cafegia", "Кав'ярня / коворкінг GIA", 175, "м²", "Бриф", NUM); r += 1
put(r, "caferent", "Оренда кав'ярні", 150, "£/м²/рік", "Припущення для невеликої локації", GBP); r += 1
put(r, "cafeyield", "Дохідність при оцінці кав'ярні", 0.075, "%", "Припущення (вторинна роздрібна торгівля)", PCT); r += 1
put(r, "apts", "Бутік-апартаменти", 4, "шт.", "Бриф", NUM); r += 1
put(r, "aptval", "Вартість апартамента (продаж)", 220000, "£/шт.", "Припущення", GBP); r += 1

inp.column_dimensions["A"].width = 42
inp.column_dimensions["B"].width = 14
inp.column_dimensions["C"].width = 14
inp.column_dimensions["D"].width = 60
for col in "EFGHIJ": inp.column_dimensions[col].width = 13
inp.column_dimensions["K"].width = 70
inp.row_dimensions[4].height = 45
inp.freeze_panes = "A5"

def T(row, col):  # ref to type table
    return f"{P}${col}${row}"


# ---------------------------------------------------------------- House cost build-up
cb = wb.create_sheet("Собівартість будинку")
cb["A1"] = "Собівартість будинку під ключ — від оболонки ZON £1 700/м²"; cb["A1"].font = TITLE
cb["A2"] = ("Оболонка ZON (£1 700/м²) містить лише будинок і систему опалення. Усе, що потрібно, щоб будинок "
            "можна було продати під іпотеку, додано нижче. Ставки — припущення центрального сценарію; діапазони в примітках.")
cb["A2"].font = Font(name=F, italic=True); cb.merge_cells("A2:F2")
cb["A2"].alignment = Alignment(wrap_text=True); cb.row_dimensions[2].height = 45
hdr(cb, 4, ["Стаття", "База", "Compact 61", "Meadow 71", "Country 140", "Примітка / діапазон"])
cb["A5"] = "GIA, м²"
for j, tr in enumerate(TYPE_ROWS):
    c = cb.cell(5, 3 + j, f"={T(tr,'B')}"); c.font = GREEN; c.number_format = NUM1
cb["A6"] = "Кількість у сценарії A"
for j in range(3):
    c = cb.cell(6, 3 + j, f"='Сценарій A'!$B${5+j}"); c.font = GREEN; c.number_format = NUM

# (label, basis, rate or per-type values, note)  basis: m2 = rate × GIA, unit = per-type £
items = [
    ("Оболонка ZON + опалення", "m2", 1700, "Від команди ZON. Стіни, дах, вікна, опалення"),
    ("Фундаменти", "m2", 160, "£120–220/м²; гвинтові палі — нижня межа"),
    ("Доставка, кран, монтаж, герметизація", "m2", 120, "£90–160/м²"),
    ("Підлоги та фарбування", "m2", 60, "£40–80/м²"),
    ("Сонячні панелі в дах (вимога Future Homes Standard)", "unit", (6500, 6500, 8500), "3–5 кВт; батарея — опція ZON Ready"),
    ("Вентиляція (dMEV)", "unit", (2000, 2000, 3500), "MVHR — +£3–4k"),
    ("Кухня з технікою", "unit", (9000, 10000, 15000), "середній клас"),
    ("Санвузли", "unit", (6000, 6000, 12000), "Country — 2 санвузли"),
    ("Гарантія 10 років, Building Control, SAP/EPC, тест повітря", "unit", (3500, 3500, 3500), "без гарантії немає іпотеки"),
    ("Роботи на ділянці: тераса, огорожа, доріжки, газон", "unit", (6000, 6000, 9000), "живоплоти з граба — тут"),
]
r = 8
section(cb, 7, "Прямі витрати на будинок (за одиницю, £)", 6)
first = r
for lab, basis, rate, note in items:
    cb.cell(r, 1, lab).font = BLACK
    if basis == "m2":
        c = cb.cell(r, 2, rate); c.font = BLUE; c.number_format = GBP
        for j in range(3):
            col = chr(ord("C") + j)
            cb.cell(r, 3 + j, f"=$B{r}*{col}$5").number_format = GBP
    else:
        cb.cell(r, 2, "за будинок").font = BLACK
        for j, v in enumerate(rate):
            c = cb.cell(r, 3 + j, v); c.font = BLUE; c.number_format = GBP
    cb.cell(r, 6, note).font = Font(name=F, italic=True, color="666666")
    r += 1
last = r - 1
cb.cell(r, 1, "Разом прямі").font = BOLD
for j in range(3):
    col = chr(ord("C") + j)
    c = cb.cell(r, 3 + j, f"=SUM({col}{first}:{col}{last})"); c.font = BOLD; c.number_format = GBP
direct = r; r += 1
cb.cell(r, 1, "Майданчик і маржа підрядника").font = BLACK
c = cb.cell(r, 2, 0.12); c.font = BLUE; c.number_format = PCT
cb.cell(r, 6, "10–15%; менше, якщо ZON сам веде будівництво").font = Font(name=F, italic=True, color="666666")
for j in range(3):
    col = chr(ord("C") + j)
    cb.cell(r, 3 + j, f"={col}{direct}*$B{r}").number_format = GBP
prel = r; r += 1
cb.cell(r, 1, "Будинок під ключ, £").font = BOLD
for j in range(3):
    col = chr(ord("C") + j)
    c = cb.cell(r, 3 + j, f"={col}{direct}+{col}{prel}"); c.font = BOLD; c.number_format = GBP
turnkey = r; r += 1
cb.cell(r, 1, "Під ключ, £/м²").font = BOLD
for j in range(3):
    col = chr(ord("C") + j)
    c = cb.cell(r, 3 + j, f"=IF({col}5>0,{col}{turnkey}/{col}5,0)"); c.font = BOLD; c.number_format = GBP
r += 2
assert r <= 27, r
cb["A27"] = "Середнє під ключ для міксу сценарію A, £/м²"; cb["A27"].font = BOLD
cb["E27"] = f"=SUMPRODUCT(C{turnkey}:E{turnkey},C6:E6)/SUMPRODUCT(C5:E5,C6:E6)"
cb["E27"].font = BOLD; cb["E27"].number_format = GBP; cb["E27"].fill = YEL
cb["A28"] = ("Не входить сюди (рахується в моделі окремо): дороги, мережі та підключення, SuDS — «Зовнішні роботи»; "
             "проєктування, резерв, S106, продаж, фінансування, земля.")
cb["A28"].font = Font(name=F, italic=True); cb.merge_cells("A28:F28")
cb["A28"].alignment = Alignment(wrap_text=True); cb.row_dimensions[28].height = 30
cb.column_dimensions["A"].width = 56; cb.column_dimensions["B"].width = 14
for col in "CDE": cb.column_dimensions[col].width = 16
cb.column_dimensions["F"].width = 46

# ---------------------------------------------------------------- Scenarios
def scenario(title, sheetname, counts, hub, route, prob, delay, abort):
    ws = wb.create_sheet(sheetname)
    ws["A1"] = title; ws["A1"].font = TITLE
    ws["A2"] = route; ws["A2"].font = Font(name=F, italic=True); ws.merge_cells("A2:F2")
    hdr(ws, 4, ["Тип юніта", "К-сть", "GIA разом, м²", "Ціна юніта, £", "GDV ринковий, £", "Ділянки, м²"])
    rows = []
    for i, (tr, n) in enumerate(zip(TYPE_ROWS, counts)):
        r = 5 + i; rows.append(r)
        ws.cell(r, 1, f"={T(tr,'A')}").font = GREEN
        c = ws.cell(r, 2, n); c.font = BLUE
        ws.cell(r, 3, f"=B{r}*{T(tr,'B')}").font = BLACK
        ws.cell(r, 4, f"={T(tr,'I')}").font = GREEN
        ws.cell(r, 5, f"=B{r}*D{r}").font = BLACK
        ws.cell(r, 6, f"=B{r}*{T(tr,'J')}").font = BLACK
    tot = 5 + len(counts)
    ws.cell(tot, 1, "Разом бунгало").font = BOLD
    for col in "BCEF":
        ws[f"{col}{tot}"] = f"=SUM({col}5:{col}{tot-1})"; ws[f"{col}{tot}"].font = BOLD
    for rr in range(5, tot + 1):
        ws[f"B{rr}"].number_format = NUM; ws[f"C{rr}"].number_format = NUM1
        ws[f"D{rr}"].number_format = GBP; ws[f"E{rr}"].number_format = GBP
        ws[f"F{rr}"].number_format = NUM
    meadow = rows[1]

    R = {}
    def line(r, key, label, formula, fmt, font=BLACK, note=None, bold=False):
        ws.cell(r, 1, label).font = BOLD if bold else BLACK
        c = ws.cell(r, 2, formula); c.font = BOLD if bold else font; c.number_format = fmt
        if note: ws.cell(r, 3, note).font = Font(name=F, italic=True, color="666666")
        R[key] = f"$B${r}"

    r = tot + 2
    section(ws, r, "Параметри сценарію", 6); r += 1
    line(r, "hub", "ZON Hub включено (1 = так, 0 = ні)", hub, NUM, BLUE); r += 1
    line(r, "prob", "Ймовірність отримати дозвіл", prob, PCT, BLUE, "Експертна оцінка, не факт"); ws[f"B{r}"].fill = YEL; r += 1
    line(r, "delay", "Додаткова затримка на планування", delay, NUM, BLUE, "місяців; фінансування землі тим часом"); r += 1
    line(r, "abort", "Витрати, якщо в дозволі відмовлять", abort, GBP, BLUE, "консультанти, звіти, апеляція"); r += 1
    line(r, "dw", "Житлових юнітів разом", f"=B{tot}+{R['hub']}*{A['apts']}", NUM, note="бунгало + апартаменти Hub"); r += 1

    r += 1
    section(ws, r, "Баланс території", 6); r += 1
    line(r, "plots", "Приватні ділянки", f"=F{tot}", NUM); r += 1
    line(r, "roads", "Дороги та паркування", f"={A['site']}*{A['roads']}", NUM, GREEN); r += 1
    line(r, "suds", "SuDS / відкритий простір", f"={A['site']}*{A['suds']}", NUM, GREEN); r += 1
    line(r, "bng", "BNG / буфери", f"={A['site']}*{A['bng']}", NUM, GREEN); r += 1
    line(r, "hubl", "ZON Hub", f"={R['hub']}*{A['hubland']}", NUM); r += 1
    line(r, "need", "Потрібно разом, м²", f"=SUM({R['plots']}:{R['hubl']})", NUM, bold=True); r += 1
    line(r, "site", "Доступно, м²", f"={A['site']}", NUM, GREEN); r += 1
    line(r, "bal", "Запас (+) / дефіцит (−), м²", f"={R['site']}-{R['need']}", NUM, bold=True); r += 1
    line(r, "fit", "Перевірка", f'=IF({R["bal"]}>=0,"Вміщується","НЕ ВМІЩУЄТЬСЯ")', "@", bold=True); r += 1
    line(r, "avgplot", "Середня ділянка, м²", f"=IF(B{tot}>0,F{tot}/B{tot},0)", NUM); r += 1

    r += 1
    section(ws, r, "Виручка (GDV)", 6); r += 1
    line(r, "omv", "Бунгало за ринковою ціною", f"=E{tot}", GBP); r += 1
    line(r, "affn", "Юнітів доступного житла (вимога)", f"=ROUNDUP({R['dw']}*{A['aff']},0)", NUM, note="округлення вгору від усіх житлових юнітів"); r += 1
    line(r, "affm", "…з них покривається обраним типом", f"=MIN({R['affn']},INDEX(B5:B{tot-1},{A['afftype']}))", NUM); r += 1
    line(r, "short", "Перевірка покриття", f'=IF({R["affm"]}<{R["affn"]},"Обраного типу замало — потрібні інші типи","OK")', "@"); r += 1
    line(r, "affd", "Знижка на доступне житло", f"=-{R['affn']}*INDEX(D5:D{tot-1},{A['afftype']})*(1-{A['affval']})", GBP, note="усі доступні юніти — за ціною обраного типу"); r += 1
    line(r, "hubv", "ZON Hub: кав'ярня (капіталізація)", f"={R['hub']}*{A['cafegia']}*{A['caferent']}/{A['cafeyield']}", GBP); r += 1
    line(r, "aptv", "ZON Hub: апартаменти", f"={R['hub']}*{A['apts']}*{A['aptval']}", GBP); r += 1
    line(r, "gdv", "GDV разом", f"=SUM({R['omv']},{R['affd']},{R['hubv']},{R['aptv']})", GBP, bold=True); r += 1

    r += 1
    section(ws, r, "Витрати (TDC)", 6); r += 1
    line(r, "build", "Будівництво бунгало", f"=C{tot}*{A['bc']}", GBP); r += 1
    line(r, "buildh", "Будівництво Hub", f"={R['hub']}*{A['hubgia']}*{A['bchub']}", GBP); r += 1
    line(r, "ext", "Зовнішні роботи та мережі", f"=({R['build']}+{R['buildh']})*{A['ext']}", GBP); r += 1
    line(r, "cont", "Резерв", f"=({R['build']}+{R['buildh']}+{R['ext']})*{A['cont']}", GBP); r += 1
    line(r, "fees", "Проєктування та консультанти", f"=({R['build']}+{R['buildh']}+{R['ext']})*{A['fees']}", GBP); r += 1
    line(r, "s106", "S106", f"={R['dw']}*{A['s106']}", GBP); r += 1
    line(r, "rams", "RAMS", f"={R['dw']}*{A['rams']}", GBP); r += 1
    line(r, "bngc", "BNG поза ділянкою", f"={A['bngoff']}", GBP, GREEN); r += 1
    line(r, "sales", "Маркетинг і агенти", f"={R['gdv']}*{A['sales']}", GBP); r += 1
    line(r, "legal", "Юридичні на продаж", f"={R['dw']}*{A['legal']}", GBP); r += 1
    line(r, "land", "Земля", f"={A['land']}", GBP, GREEN); r += 1
    line(r, "landc", "SDLT + юридичні на землю", f"={R['land']}*{A['landcost']}", GBP); r += 1
    line(r, "fin", "Фінансування",
         f"={A['rate']}*({A['draw']}*SUM({R['build']}:{R['legal']})*{A['months']}/12"
         f"+({R['land']}+{R['landc']})*({A['months']}+{R['delay']})/12)", GBP,
         note="будівництво: середня вибірка; земля: весь строк + затримка"); r += 1
    line(r, "tdc", "TDC разом", f"=SUM({R['build']}:{R['fin']})", GBP, bold=True); r += 1

    r += 1
    section(ws, r, "Результат", 6); r += 1
    line(r, "profit", "Прибуток девелопера", f"={R['gdv']}-{R['tdc']}", GBP, bold=True); r += 1
    line(r, "mgdv", "Маржа на GDV", f"=IF({R['gdv']}<>0,{R['profit']}/{R['gdv']},0)", PCT, bold=True); r += 1
    line(r, "mcost", "Маржа на витрати", f"=IF({R['tdc']}<>0,{R['profit']}/{R['tdc']},0)", PCT); r += 1
    line(r, "ok", "Проти цільової маржі", f'=IF({R["mgdv"]}>={A["target"]},"Досягає цілі","НИЖЧЕ ЦІЛІ")', "@", bold=True); r += 1
    line(r, "rlv", "Залишкова вартість землі при цільовій маржі",
         f"=({R['gdv']}*(1-{A['target']})-({R['tdc']}-{R['land']}-{R['landc']}"
         f"-{A['rate']}*({R['land']}+{R['landc']})*({A['months']}+{R['delay']})/12))"
         f"/(1+{A['landcost']}+{A['rate']}*(1+{A['landcost']})*({A['months']}+{R['delay']})/12)",
         GBP, note="скільки максимум можна заплатити за землю"); r += 1
    line(r, "risk", "Прибуток з поправкою на ризик планування",
         f"={R['prob']}*{R['profit']}-(1-{R['prob']})*{R['abort']}", GBP, bold=True); r += 1
    r += 1
    section(ws, r, "Для аналізу чутливості", 6); r += 1
    line(r, "blink", "Витрати, що залежать від будівництва",
         f"=SUM({R['build']}:{R['fees']})", GBP); r += 1
    line(r, "fixed", "Інші витрати (без маркетингу)",
         f"={R['tdc']}-{R['blink']}-{R['sales']}", GBP); r += 1
    line(r, "be", "Потрібна зміна цін для цільової маржі",
         f"=({R['blink']}+{R['fixed']})/({R['gdv']}*(1-{A['sales']}-{A['target']}))-1", PCT,
         note="при поточній ціні землі; фінансування — константа"); r += 1

    ws.column_dimensions["A"].width = 46
    for col in "BCDEF": ws.column_dimensions[col].width = 18
    ws.row_dimensions[4].height = 32
    return ws, R

sa, RA = scenario("Сценарій A — 25 бунгало в межах outline consent", "Сценарій A",
                  [10, 8, 7], 0,
                  "Шлях: Reserved Matters. Без Hub. Відповідає ліміту 25 юнітів (за словами користувача).",
                  0.85, 0, 100000)
sb, RB = scenario("Сценарій B — бриф: 25 бунгало + ZON Hub (4 апартаменти + кав'ярня) = 29 житлових", "Сценарій B",
                  [10, 8, 7], 1,
                  "Шлях: бунгало в межах 25, але Hub додає 4 житла, клас E і короткострокову оренду — окрема / нова заявка. Ризик відмови.",
                  0.45, 12, 250000)

# ---------------------------------------------------------------- Summary
sm = wb.create_sheet("Порівняння", 0)
sm["A1"] = "ZON Wivenhoe — порівняння сценаріїв"; sm["A1"].font = TITLE
sm["A2"] = ("Усі показники рахуються з аркушів сценаріїв. Ціни, вартість будівництва і земля — "
            "ПРИПУЩЕННЯ (див. «Припущення»), а не підтверджені дані.")
sm["A2"].font = Font(name=F, italic=True); sm.merge_cells("A2:D2")
hdr(sm, 4, ["Показник", "A: 25 бунгало (RM)", "B: бриф 25 + Hub", "Різниця B − A"])
items = [
    ("Житлових юнітів", "dw", NUM), ("Баланс території, м²", "bal", NUM),
    ("Перевірка території", "fit", "@"), ("Середня ділянка, м²", "avgplot", NUM),
    ("Доступних юнітів", "affn", NUM), ("GDV", "gdv", GBP), ("TDC", "tdc", GBP),
    ("Прибуток", "profit", GBP), ("Маржа на GDV", "mgdv", PCT),
    ("Проти цілі", "ok", "@"), ("Макс. ціна землі (RLV)", "rlv", GBP),
    ("Потрібна зміна цін для цілі", "be", PCT),
    ("Ймовірність дозволу", "prob", PCT), ("Прибуток з поправкою на ризик", "risk", GBP),
]
for i, (lab, key, fmt) in enumerate(items):
    r = 5 + i
    sm.cell(r, 1, lab).font = BLACK
    a = sm.cell(r, 2, f"='Сценарій A'!{RA[key]}"); a.font = GREEN; a.number_format = fmt
    b = sm.cell(r, 3, f"='Сценарій B'!{RB[key]}"); b.font = GREEN; b.number_format = fmt
    if fmt != "@":
        d = sm.cell(r, 4, f"=C{r}-B{r}"); d.font = BLACK; d.number_format = fmt
for key in ("profit", "risk", "mgdv"):
    idx = [k for _, k, _ in items].index(key) + 5
    for col in "ABCD": sm[f"{col}{idx}"].font = Font(name=F, bold=True, color=sm[f"{col}{idx}"].font.color)

r0 = 5 + len(items) + 2
sens = [-0.10, -0.05, 0, 0.05, 0.10]
for s_i, (nm, R) in enumerate([("Сценарій A", RA), ("Сценарій B", RB)]):
    top = r0 + s_i * 9
    sm.cell(top, 1, f"Чутливість прибутку — {nm} (фінансування — константа)").font = BOLD
    sm.cell(top + 1, 1, "Ціна продажу ↓ / Вартість будівництва →").font = Font(name=F, italic=True)
    for j, dc in enumerate(sens):
        c = sm.cell(top + 1, 2 + j, dc); c.font = BLUE; c.number_format = '+0%;-0%;0%'
    for i, dp in enumerate(sens):
        rr = top + 2 + i
        c = sm.cell(rr, 1, dp); c.font = BLUE; c.number_format = '+0%;-0%;0%'
        for j in range(len(sens)):
            col = chr(ord("B") + j)
            q = f"'{nm}'!"
            f = (f"={q}{R['gdv']}*(1+$A{rr})*(1-{A['sales']})"
                 f"-{q}{R['blink']}*(1+{col}${top+1})-{q}{R['fixed']}")
            c = sm.cell(rr, 2 + j, f); c.font = BLACK; c.number_format = GBP
sm.column_dimensions["A"].width = 44
for col in "BCDEF": sm.column_dimensions[col].width = 20
sm.row_dimensions[4].height = 32


# ---------------------------------------------------------------- Improved case
im = wb.create_sheet("Покращений випадок", 1)
im["A1"] = "Improved Case — важелі поверх сценарію A (25 бунгало)"; im["A1"].font = TITLE
im["A2"] = ("Кожен важіль рахується окремо від бази сценарію A. Сині клітинки — припущення важелів "
            "(з запропонованої стратегії), жовті — ті, що найбільше впливають і потребують підтвердження.")
im["A2"].font = Font(name=F, italic=True); im.merge_cells("A2:D2")
im["A2"].alignment = Alignment(wrap_text=True); im.row_dimensions[2].height = 30
Q = "'Сценарій A'!"
I = {}
def il(r, key, label, val, fmt, font=BLACK, note=None, bold=False, yellow=False):
    im.cell(r, 1, label).font = BOLD if bold else BLACK
    c = im.cell(r, 2, val); c.font = BOLD if bold else font; c.number_format = fmt
    if yellow: c.fill = YEL
    if note: im.cell(r, 3, note).font = Font(name=F, italic=True, color="666666")
    I[key] = f"$B${r}"

r = 4
section(im, r, "База: сценарій A", 4); r += 1
il(r, "gdv0", "GDV бази", f"={Q}{RA['gdv']}", GBP, GREEN); r += 1
il(r, "p0", "Прибуток бази", f"={Q}{RA['profit']}", GBP, GREEN); r += 1
il(r, "m0", "Маржа бази", f"={Q}{RA['mgdv']}", PCT, GREEN); r += 1

r += 1
section(im, r, "1. ZON Hub — продаж за капіталізованою вартістю (без апартаментів)", 4); r += 1
il(r, "hinc", "Чистий річний дохід Hub", 27500, GBP, BLUE, "стратегія: £25–30k; потрібен договір з оператором", yellow=True); r += 1
il(r, "hyld", "Ставка капіталізації", 0.07, PCT, BLUE, "стратегія: 7%; для кав'ярні реалістичніше 8–10%", yellow=True); r += 1
il(r, "hcost", "Повна вартість Hub (будівля + fit-out)", 390000, GBP, BLUE, "стратегія: £390k ≈ £2 230/м² на 175 м² — перевірити QS", yellow=True); r += 1
il(r, "hval", "Капіталізована вартість", f"={I['hinc']}/{I['hyld']}", GBP); r += 1
il(r, "hup", "Приріст прибутку від Hub", f"={I['hval']}*(1-{A['sales']})-{I['hcost']}", GBP, bold=True,
   note="у сценарії A Hub не закладено, тож віднімаємо і його вартість"); r += 1

r += 1
section(im, r, "2. Пакети ZON Ready", 4); r += 1
il(r, "zmkt", "Ринкових юнітів (без доступного житла)", f"={Q}{RA['dw']}-{Q}{RA['affn']}", NUM,
   note="RSL не купує апгрейди"); r += 1
il(r, "zupt", "Частка покупців, що беруть пакет", 0.6, PCT, BLUE, "стратегія припускала 100% на 25 юнітів", yellow=True); r += 1
il(r, "zprice", "Ціна пакета", 35000, GBP, BLUE, "стратегія"); r += 1
il(r, "zcost", "Собівартість пакета", 20000, GBP, BLUE, "стратегія"); r += 1
il(r, "zn", "Пакетів продано", f"=ROUND({I['zmkt']}*{I['zupt']},0)", NUM); r += 1
il(r, "zrev", "Додаткова виручка", f"={I['zn']}*{I['zprice']}", GBP); r += 1
il(r, "zup", "Приріст прибутку від пакетів", f"={I['zn']}*({I['zprice']}*(1-{A['sales']})-{I['zcost']})", GBP, bold=True); r += 1

r += 1
section(im, r, "3. Value engineering", 4); r += 1
il(r, "ve", "Економія на будівельних витратах", 0.04, PCT, BLUE, "стратегія: 3–5%; підтверджує QS"); r += 1
il(r, "veup", "Приріст прибутку", f"={Q}{RA['blink']}*{I['ve']}", GBP, bold=True); r += 1

r += 1
section(im, r, "4. Переговори про ціну землі", 4); r += 1
il(r, "ld", "Знижка на землю", 140000, GBP, BLUE, "стратегія: £100–180k"); r += 1
il(r, "ldup", "Приріст прибутку (з SDLT і фінансуванням землі)",
   f"={I['ld']}*(1+{A['landcost']})*(1+{A['rate']}*({A['months']}+{Q}{RA['delay']})/12)", GBP, bold=True); r += 1

r += 1
section(im, r, "5. Зниження витрат на продаж", 4); r += 1
il(r, "sc", "Зниження статті «маркетинг і агенти»", 0.25, PCT, BLUE, "стратегія: £50–70k"); r += 1
il(r, "scup", "Приріст прибутку", f"={Q}{RA['sales']}*{I['sc']}", GBP, bold=True); r += 1

r += 1
section(im, r, "6. Траншеве залучення equity (довідково, НЕ в сумі)", 4); r += 1
il(r, "eq", "Equity разом", 2560000, GBP, BLUE, "стратегія"); r += 1
il(r, "eqr", "Priority return", 0.12, PCT, BLUE, "стратегія"); r += 1
il(r, "eq1", "Перший транш", 800000, GBP, BLUE, "стратегія"); r += 1
il(r, "eqm", "Строк, місяців", 24, NUM, BLUE); r += 1
il(r, "eqm2", "Строк другого траншу, місяців", 15, NUM, BLUE); r += 1
il(r, "eqfull", "Return при внесенні всього одразу", f"={I['eq']}*{I['eqr']}*{I['eqm']}/12", GBP); r += 1
il(r, "eqtr", "Return при траншах",
   f"={I['eq1']}*{I['eqr']}*{I['eqm']}/12+({I['eq']}-{I['eq1']})*{I['eqr']}*{I['eqm2']}/12", GBP); r += 1
il(r, "eqsave", "Економія", f"={I['eqfull']}-{I['eqtr']}", GBP, bold=True,
   note="база моделі фінансується боргом 9%, а не equity 12% — тому не додаємо до суми"); r += 1

r += 1
section(im, r, "Результат Improved Case", 4); r += 1
il(r, "up", "Сума важелів 1–5", f"={I['hup']}+{I['zup']}+{I['veup']}+{I['ldup']}+{I['scup']}", GBP, bold=True); r += 1
il(r, "p1", "Прибуток Improved Case", f"={I['p0']}+{I['up']}", GBP, bold=True); r += 1
il(r, "g1", "GDV Improved Case", f"={I['gdv0']}+{I['hval']}+{I['zrev']}", GBP); r += 1
il(r, "m1", "Маржа на GDV", f"=IF({I['g1']}<>0,{I['p1']}/{I['g1']},0)", PCT, bold=True); r += 1
il(r, "gap", "Скільки ще бракує до цільової маржі", f"=MAX(0,{A['target']}*{I['g1']}-{I['p1']})", GBP, bold=True); r += 1
im.column_dimensions["A"].width = 52; im.column_dimensions["B"].width = 18; im.column_dimensions["C"].width = 70

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is not None and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic,
                              color=c.font.color, size=c.font.size)
wb.save(OUT)
print("saved", OUT)
